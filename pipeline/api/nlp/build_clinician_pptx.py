#!/usr/bin/env python3
"""Build a NATIVE, editable PowerPoint from docs/reports/clinician_overview.md.
Editable text + native tables + embedded chart images. Parses the Marp markdown so
it stays in sync with the rendered deck.
"""
import html
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path("/home/zulaika/ClinicalAnnotationWeb")
MD = ROOT / "docs" / "reports" / "clinician_overview.md"
OUT = ROOT / "docs" / "reports" / "clinician_overview_editable.pptx"
MDDIR = MD.parent

BLUE = RGBColor(0x1A, 0x4F, 0x8B); GREEN = RGBColor(0x1A, 0x7F, 0x37); RED = RGBColor(0xB4, 0x23, 0x18)
DARK = RGBColor(0x14, 0x18, 0x1F); GREY = RGBColor(0x55, 0x55, 0x55); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHTBLUE = RGBColor(0xF3, 0xF6, 0xFB)
EMW, EMH = Inches(13.333), Inches(7.5)


def clean(s):
    s = re.sub(r"</?span[^>]*>", "", s)  # drop span tags (color handled separately)
    s = re.sub(r"<[^>]+>", "", s)
    return html.unescape(s).strip()


SPAN = re.compile(r'<span class="(good|bad)">(.*?)</span>|\*\*(.*?)\*\*|\*(.*?)\*')


def runs_of(text):
    """-> list of (text, bold, italic, color)."""
    out, pos = [], 0
    for m in SPAN.finditer(text):
        if m.start() > pos:
            t = clean(text[pos:m.start()])
            if t:
                out.append((t, False, False, None))
        if m.group(1):
            out.append((clean(m.group(2)), True, False, GREEN if m.group(1) == "good" else RED))
        elif m.group(3) is not None:
            out.append((clean(m.group(3)), True, False, None))
        elif m.group(4) is not None:
            out.append((clean(m.group(4)), False, True, None))
        pos = m.end()
    if pos < len(text):
        t = clean(text[pos:])
        if t:
            out.append((t, False, False, None))
    return out or [("", False, False, None)]


def set_runs(paragraph, text, size, base_color=DARK, base_bold=False):
    paragraph.clear() if hasattr(paragraph, "clear") else None
    for i, (t, b, it, col) in enumerate(runs_of(text)):
        r = paragraph.add_run()
        r.text = (" " if i and not t.startswith(" ") else "") + t
        r.font.size = Pt(size); r.font.bold = b or base_bold
        r.font.italic = it
        r.font.color.rgb = col or base_color


def split_slides(md):
    body = md.split("---", 2)[2]                 # drop YAML front matter (first two ---)
    chunks, cur = [], []
    for line in body.split("\n"):
        if line.strip() == "---":
            chunks.append(cur); cur = []
        else:
            cur.append(line)
    chunks.append(cur)
    return [c for c in chunks if any(x.strip() for x in c)]


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    return tf


def main():
    prs = Presentation(); prs.slide_width = EMW; prs.slide_height = EMH
    blank = prs.slide_layouts[6]
    md = MD.read_text(encoding="utf-8")

    for chunk in split_slides(md):
        lines = chunk
        lead = any("_class: lead" in ln for ln in lines)
        slide = prs.slides.add_slide(blank)
        # collect blocks
        i = 0
        title = subtitle = note = None
        blocks = []                              # ('bullets',[(text,level)]) ('table',rows) ('image',path) ('quote',text) ('big',text) ('para',text)
        bullets, table = [], []
        while i < len(lines):
            ln = lines[i]; s = ln.strip()
            if not s or s.startswith("<!--"):
                i += 1; continue
            if s.startswith("### "):
                subtitle = clean(s[4:]); i += 1; continue
            if s.startswith("## "):
                title = clean(s[3:]); i += 1; continue
            if s.startswith("# "):
                title = clean(s[2:]); i += 1; continue
            if 'class="note"' in s:
                note = clean(s); i += 1; continue
            if s.startswith(">"):
                blocks.append(("quote", clean(s[1:]))); i += 1; continue
            if s.startswith("|"):
                row = [c.strip() for c in s.strip("|").split("|")]
                if not set("".join(row)) <= set("-: "):
                    table.append(row)
                if not (i + 1 < len(lines) and lines[i + 1].strip().startswith("|")):
                    if table:
                        blocks.append(("table", table)); table = []
                i += 1; continue
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
            if m:
                blocks.append(("image", m.group(1))); i += 1; continue
            if 'class="big"' in s or s.startswith("<div"):
                i += 1; continue
            if s.startswith("</div>"):
                i += 1; continue
            mb = re.match(r"^(\s*)([-\d.]+)\s+(.*)", ln)
            if mb and (mb.group(2).startswith("-") or mb.group(2)[0].isdigit()):
                lvl = len(mb.group(1)) // 2
                bullets.append((mb.group(3), lvl)); i += 1; continue
            blocks.append(("para", s)); i += 1
        if bullets:
            blocks.insert(0, ("bullets", bullets))

        # ---- render ----
        if lead:
            tf = add_textbox(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.6), MSO_ANCHOR.MIDDLE)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            set_runs(p, title or "", 40, BLUE, True)
            if subtitle:
                p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
                set_runs(p2, subtitle, 24, DARK)
            for typ, val in blocks:
                if typ in ("para", "quote"):
                    pp = tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER
                    set_runs(pp, val, 18, GREY)
            continue

        # title
        if title:
            tf = add_textbox(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(1.0))
            set_runs(tf.paragraphs[0], title, 30, BLUE, True)
        top = Inches(1.55)
        for typ, val in blocks:
            if typ == "bullets":
                tf = add_textbox(slide, Inches(0.7), top, Inches(12.0), Inches(5.0))
                first = True
                for text, lvl in val:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p.level = min(lvl, 4); p.space_after = Pt(8)
                    bullet = ("- " if lvl else "") if False else ""
                    set_runs(p, text, 20 - 2 * min(lvl, 2))
                top = Inches(6.8)
            elif typ == "table":
                rows, cols = len(val), max(len(r) for r in val)
                h = Inches(min(0.55 * rows + 0.2, 4.8))
                gt = slide.shapes.add_table(rows, cols, Inches(0.6), top, Inches(12.1), h).table
                for ci in range(cols):
                    gt.columns[ci].width = Inches(12.1 / cols)
                for ri, r in enumerate(val):
                    for ci in range(cols):
                        cell = gt.cell(ri, ci)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        tcell = cell.text_frame; tcell.word_wrap = True
                        txt = r[ci] if ci < len(r) else ""
                        if ri == 0:
                            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
                            set_runs(tcell.paragraphs[0], txt, 14, WHITE, True)
                        else:
                            set_runs(tcell.paragraphs[0], txt, 13)
                top = top + h + Inches(0.15)
            elif typ == "image":
                p = (MDDIR / val).resolve()
                if p.exists():
                    pic = slide.shapes.add_picture(str(p), Inches(2.4), top, width=Inches(8.5))
                    # center horizontally
                    pic.left = int((EMW - pic.width) / 2)
                    top = top + pic.height + Inches(0.1)
            elif typ == "quote":
                sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), top, Inches(12.1), Inches(1.0))
                sp.fill.solid(); sp.fill.fore_color.rgb = LIGHTBLUE
                sp.line.color.rgb = BLUE; sp.line.width = Pt(1.5)
                sp.shadow.inherit = False
                tf = sp.text_frame; tf.word_wrap = True
                tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
                set_runs(tf.paragraphs[0], val, 16, DARK)
                top = top + Inches(1.2)
            elif typ == "para":
                tf = add_textbox(slide, Inches(0.7), top, Inches(12.0), Inches(1.0))
                set_runs(tf.paragraphs[0], val, 19)
                top = top + Inches(0.7)
        if note:
            tf = add_textbox(slide, Inches(0.6), Inches(6.85), Inches(12.1), Inches(0.5))
            set_runs(tf.paragraphs[0], note, 13, GREY)

    prs.save(str(OUT))
    print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
